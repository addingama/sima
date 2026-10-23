#!/usr/bin/env python3
"""Deck berurutan untuk setup demo staging dan uji coba tim.

  python3 -m venv .venv-pptx
  .venv-pptx/bin/pip install python-pptx
  .venv-pptx/bin/python docs/demo/generate_uji_coba_staging_pptx.py
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree
from pathlib import Path

OUT = Path(__file__).resolve().parent / "SIMA-Uji-Coba-Staging.pptx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_MID = RGBColor(0x14, 0x32, 0x54)
TEAL = RGBColor(0x0D, 0x7A, 0x6F)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
SLATE = RGBColor(0x2D, 0x3A, 0x4A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF7, 0xF5, 0xF2)
LIGHT = RGBColor(0xEE, 0xF3, 0xF6)
AMBER_BG = RGBColor(0xFB, 0xF3, 0xE4)
ROSE_BG = RGBColor(0xFB, 0xEE, 0xEE)
ROSE = RGBColor(0x9F, 0x2D, 0x2D)

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 24


def set_run(run, size=18, bold=False, color=SLATE, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_rect(slide, left, top, width, height, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.08
    return shape


def send_to_back(slide, shape):
    tree = slide.shapes._spTree
    element = shape._element
    tree.remove(element)
    tree.insert(2, element)


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=SLATE, align=PP_ALIGN.LEFT, name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    set_run(paragraph.runs[0], size=size, bold=bold, color=color, name=name)
    return box


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_bg(slide, color=OFF_WHITE):
    send_to_back(slide, add_rect(slide, 0, 0, W, H, color))


def add_header_bar(slide):
    add_rect(slide, 0, 0, W, Inches(0.08), TEAL)
    add_rect(slide, 0, Inches(7.42), W, Inches(0.08), NAVY)


def add_footer(slide, page):
    add_textbox(
        slide, Inches(0.55), Inches(7.08), Inches(9.4), Inches(0.28),
        "SIMA  ·  Uji coba staging  ·  23 Oktober 2026",
        size=11, color=MUTED,
    )
    add_textbox(
        slide, Inches(10.4), Inches(7.08), Inches(2.35), Inches(0.28),
        f"{page:02d}  /  {TOTAL:02d}",
        size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def content_title(slide, title, kicker):
    add_header_bar(slide)
    add_textbox(slide, Inches(0.55), Inches(0.24), Inches(12.2), Inches(0.28), kicker.upper(), size=12, bold=True, color=TEAL)
    add_textbox(slide, Inches(0.55), Inches(0.5), Inches(12.2), Inches(0.5), title, size=28, bold=True, color=NAVY)


def bullets(slide, lines, top, size=16, color=SLATE, left=0.55, width=12.2, gap=0.42):
    for index, line in enumerate(lines):
        add_textbox(
            slide, Inches(left), Inches(top + index * gap), Inches(width), Inches(gap),
            line, size=size, color=color,
        )


def _set_cell_border(cell, color="D6DEE6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        line = etree.SubElement(tc_pr, qn(f"a:{edge}"), w="6350", cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(line, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=color)


def add_table(slide, headers, rows, left, top, width, height, font=13, widths=None):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = shape.table
    if widths:
        for index, column_width in enumerate(widths):
            table.columns[index].width = column_width
    else:
        each = int(width / len(headers))
        for index in range(len(headers)):
            table.columns[index].width = each
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                set_run(run, size=font, bold=True, color=WHITE)
        _set_cell_border(cell, "0B1F3A")
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index + 1, col_index)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_index % 2 == 0 else LIGHT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    set_run(run, size=font, color=SLATE)
            _set_cell_border(cell)
    return shape


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def title_slide(prs):
    slide = blank(prs)
    add_bg(slide, NAVY)
    add_rect(slide, 0, 0, Inches(0.22), H, TEAL)
    add_rect(slide, 0, Inches(6.45), W, Inches(1.05), NAVY_MID)
    add_textbox(slide, Inches(0.85), Inches(1.35), Inches(11.5), Inches(0.32), "SESI HARI INI  ·  23 OKTOBER 2026", size=14, bold=True, color=GOLD)
    add_textbox(slide, Inches(0.85), Inches(1.85), Inches(11.5), Inches(1.5), "Uji coba SIMA\ndi server staging", size=44, bold=True, color=WHITE)
    add_textbox(
        slide, Inches(0.85), Inches(4.15), Inches(11), Inches(1.1),
        "Siapkan data demo, lalu setiap peran mencoba alurnya sendiri.\nIkuti urutan slide. Jangan loncat ke transaksi sebelum data demo siap.",
        size=18, color=RGBColor(0xCB, 0xD5, 0xE1),
    )
    add_textbox(slide, Inches(0.85), Inches(6.72), Inches(11.5), Inches(0.4), "https://sima.ahpeduli.org", size=16, color=WHITE)
    add_notes(slide, "Buka dengan tujuan: hari ini staging dipakai bersama, bukan go-live produksi. Satu orang memegang satu akun.")


def rules_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Tiga aturan sebelum mulai", "Berlaku untuk semua orang")
    cards = [
        ("01", "Ini staging, bukan produksi", "Alamatnya sima.ahpeduli.org. Jangan masukkan data donatur, rekening, atau nominal sungguhan."),
        ("02", "Satu orang, satu email", "Password sementara semua pengurus: @Password123. Setelah masuk, ganti lewat menu akun. Jangan berbagi password baru."),
        ("03", "Jangan hapus transaksi", "Yang sudah disetujui hanya bisa dibatalkan lewat reversal. Salah input sebelum approve masih bisa ditolak."),
    ]
    for index, (num, title, body) in enumerate(cards):
        top = Inches(1.4 + index * 1.75)
        add_rect(slide, Inches(0.55), top, Inches(12.2), Inches(1.58), WHITE, rounded=True)
        add_rect(slide, Inches(0.55), top, Inches(0.12), Inches(1.58), TEAL)
        add_textbox(slide, Inches(1.0), top + Inches(0.28), Inches(1.1), Inches(0.5), num, size=26, bold=True, color=TEAL)
        add_textbox(slide, Inches(2.3), top + Inches(0.22), Inches(9.9), Inches(0.42), title, size=20, bold=True, color=NAVY)
        add_textbox(slide, Inches(2.3), top + Inches(0.72), Inches(9.9), Inches(0.65), body, size=16, color=SLATE)
    add_footer(slide, 2)
    add_notes(slide, "Password sementara sama supaya mudah diucapkan di ruang. Begitu masuk, masing-masing mengganti password sendiri.")


def url_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Masuk ke sistem", "Langkah bersama, setelah data demo siap")
    add_rect(slide, Inches(0.55), Inches(1.45), Inches(12.2), Inches(2.15), NAVY, rounded=True)
    add_textbox(slide, Inches(0.9), Inches(1.65), Inches(11.5), Inches(0.35), "ALAMAT", size=12, bold=True, color=GOLD)
    add_textbox(slide, Inches(0.9), Inches(2.05), Inches(11.5), Inches(0.55), "https://sima.ahpeduli.org", size=32, bold=True, color=WHITE)
    add_textbox(slide, Inches(0.9), Inches(2.75), Inches(11.5), Inches(0.5), "Halaman masuk: /auth/v2/login    ·    password sementara pengurus: @Password123", size=16, color=RGBColor(0xCB, 0xD5, 0xE1))
    bullets(slide, [
        "Buka dari laptop masing-masing. Masuk dengan email di daftar pengurus, bukan email @sima.test.",
        "Karakter pertama password adalah tanda @. Huruf P besar.",
        "Menu yang terlihat berbeda-beda. Itu sesuai peran SIMA, bukan error.",
    ], top=4.0, size=17, gap=0.55)
    add_footer(slide, 3)


def accounts_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Akun pengurus", "Password sementara semua baris: @Password123")
    add_table(
        slide,
        ["Nama", "Jabatan", "Peran di SIMA", "Email"],
        [
            ["Muharrar", "Pengawas", "Auditor", "mhrrkdi@gmail.com"],
            ["Johan Saputra", "Ketua Yayasan", "Ketua", "hun7m3@gmail.com"],
            ["Syarif Ahmadi", "Sekretaris", "Verifikator", "ahmadisy@gmail.com"],
            ["Novianto", "Bendahara", "Bendahara", "novian1980@gmail.com"],
            ["Firman Ardiansyah", "Ketua LKS", "Ketua", "firman.abdullah@gmail.com"],
            ["Erwin Ahadi", "Bendahara LKS", "Asisten bendahara", "erwin.ahadi74@gmail.com"],
        ],
        Inches(0.55), Inches(1.32), Inches(12.2), Inches(4.15),
        font=13,
        widths=[Inches(2.55), Inches(2.35), Inches(2.55), Inches(4.75)],
    )
    add_textbox(
        slide, Inches(0.55), Inches(5.6), Inches(12.2), Inches(1.2),
        "Fasilitator membuat keenam akun di Pengaturan setelah data demo terisi. Peran SIMA mengikuti kolom ini.\n"
        "Portal donatur tetap akun contoh donatur@sima.test / password. Akun admin@sima.test hanya untuk yang menyiapkan server.",
        size=15, color=SLATE,
    )
    add_footer(slide, 4)
    add_notes(
        slide,
        "Dua peran Ketua: Johan mengurus persetujuan uang yayasan, Firman mengurus pengajuan bantuan. "
        "Erwin menginput karena Bendahara LKS menjalankan operasional harian; Novianto yang menyetujui. "
        "Tidak ada petugas lapangan di daftar, jadi bantuan dibuat oleh Firman.",
    )


def agenda_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Urutan sesi", "Jangan mulai bagian tim sebelum IT selesai")
    items = [
        ("A", "IT menyiapkan staging", "Backup, kosongkan data lama jika boleh, isi data demo, cek health."),
        ("B", "Semua orang masuk", "Satu akun per orang. Lihat menu yang muncul."),
        ("C", "Keuangan berurutan", "Penerimaan, lalu pengeluaran, lalu kas kecil."),
        ("D", "Bantuan dan portal", "Firman dan Johan di Kanban. Portal memakai akun donatur contoh."),
        ("E", "Cek bersama", "Auditor memastikan selisih nol. Tim mencatat temuan."),
    ]
    for index, (num, title, body) in enumerate(items):
        top = Inches(1.35 + index * 1.05)
        add_rect(slide, Inches(0.55), top, Inches(12.2), Inches(0.92), WHITE, rounded=True)
        add_textbox(slide, Inches(0.8), top + Inches(0.2), Inches(0.7), Inches(0.5), num, size=22, bold=True, color=TEAL)
        add_textbox(slide, Inches(1.7), top + Inches(0.1), Inches(10.6), Inches(0.36), title, size=18, bold=True, color=NAVY)
        add_textbox(slide, Inches(1.7), top + Inches(0.46), Inches(10.6), Inches(0.36), body, size=14, color=MUTED)
    add_footer(slide, 5)


def it_backup_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Backup dulu", "Bagian IT  ·  hanya di host staging")
    add_rect(slide, Inches(0.55), Inches(1.4), Inches(12.2), Inches(1.15), ROSE_BG, rounded=True)
    add_textbox(
        slide, Inches(0.85), Inches(1.58), Inches(11.6), Inches(0.8),
        "Perintah di bawah mengosongkan database staging. Jangan dijalankan di server produksi. Pastikan Anda berada di /opt/sima-ahp pada VPS staging.",
        size=16, color=ROSE,
    )
    bullets(slide, [
        "1.  Masuk SSH ke VPS staging, lalu cd /opt/sima-ahp.",
        "2.  Cek health: curl -s https://sima.ahpeduli.org/api/health",
        "3.  Backup database sebelum apa pun dihapus:",
    ], top=2.8, size=16, gap=0.42)
    add_rect(slide, Inches(0.55), Inches(4.2), Inches(12.2), Inches(1.15), NAVY, rounded=True)
    add_textbox(
        slide, Inches(0.8), Inches(4.4), Inches(11.7), Inches(0.8),
        "docker compose -f docker-compose.prod.yml -f docker-compose.staging.yml exec app php artisan sima:backup-db",
        size=14, color=WHITE, name="Consolas",
    )
    add_textbox(slide, Inches(0.55), Inches(5.55), Inches(12.2), Inches(1.2), "Simpan nama file backup. Jika uji coba merusak angka, restore dari file itu. Jangan lanjut jika health gagal.", size=16, color=SLATE)
    add_footer(slide, 6)
    add_notes(slide, "Compose production mengunci APP_ENV=production. Staging memakai file yang sama plus docker-compose.staging.yml.")


def it_seed_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Isi data demo", "Bagian IT  ·  dua perintah, berurutan")
    add_textbox(slide, Inches(0.55), Inches(1.25), Inches(12.2), Inches(0.4), "1. Kosongkan dan isi role, dana sistem, serta akun contoh", size=16, bold=True, color=NAVY)
    add_rect(slide, Inches(0.55), Inches(1.7), Inches(12.2), Inches(0.85), NAVY, rounded=True)
    add_textbox(
        slide, Inches(0.75), Inches(1.88), Inches(11.8), Inches(0.55),
        "docker compose -f docker-compose.prod.yml -f docker-compose.staging.yml exec app php artisan migrate:fresh --seed --force",
        size=13, color=WHITE, name="Consolas",
    )
    add_textbox(slide, Inches(0.55), Inches(2.7), Inches(12.2), Inches(0.7), "2. Isi master, saldo awal, antrian approval, kas kecil, dan donatur. Container menolak perintah ini selama terbaca sebagai production, jadi environment proses ini saja diubah ke staging, lalu cache dikembalikan.", size=15, color=SLATE)
    commands = [
        "docker compose -f docker-compose.prod.yml -f docker-compose.staging.yml exec app php artisan config:clear",
        "docker compose -f docker-compose.prod.yml -f docker-compose.staging.yml exec -e APP_ENV=staging app php artisan sima:seed-demo --force",
        "docker compose -f docker-compose.prod.yml -f docker-compose.staging.yml exec app php artisan config:cache",
    ]
    for index, command in enumerate(commands):
        top = 3.55 + index * 0.95
        add_rect(slide, Inches(0.55), Inches(top), Inches(12.2), Inches(0.82), WHITE, rounded=True)
        add_textbox(slide, Inches(0.75), Inches(top + 0.16), Inches(11.8), Inches(0.52), command, size=12, color=NAVY, name="Consolas")
    add_footer(slide, 7)
    add_notes(slide, "Jangan jalankan config:cache di dalam proses yang APP_ENV-nya staging. Perintah ketiga mengembalikan cache production. Jika muncul akun DEMO-KAS sudah ada, ulangi dari migrate:fresh.")


def it_check_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Cek sebelum tim masuk", "Bagian IT  ·  lima tanda data demo hidup")
    rows = [
        ["Health", "https://sima.ahpeduli.org/api/health menjawab database dan cache terhubung"],
        ["Akun pengurus", "Enam email pengurus sudah dibuat, password sementara @Password123"],
        ["Master demo", "Ada DEMO-BCA, DEMO-KAS, DEMO-ZKT, DEMO-YATIM, DEMO-UMUM"],
        ["Antrian", "Approval punya penerimaan submitted, pengeluaran submitted, dan satu verified"],
        ["Donatur contoh", "donatur@sima.test / password hanya melihat donasi Budi Santoso"],
    ]
    for index, (title, body) in enumerate(rows):
        top = Inches(1.35 + index * 1.02)
        add_rect(slide, Inches(0.55), top, Inches(12.2), Inches(0.9), WHITE, rounded=True)
        add_textbox(slide, Inches(0.85), top + Inches(0.22), Inches(2.3), Inches(0.45), title, size=16, bold=True, color=TEAL)
        add_textbox(slide, Inches(3.2), top + Inches(0.22), Inches(9.2), Inches(0.5), body, size=15, color=SLATE)
    add_footer(slide, 8)
    add_notes(slide, "Kas kecil: petugas.bantuan adalah PJ. Ada float yang masih menunggu cair dan ada yang sedang dilaporkan.")


def concept_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Yang harus diingat saat mencoba", "Dua angka yang selalu berpasangan")
    add_rect(slide, Inches(0.55), Inches(1.4), Inches(6.0), Inches(3.5), WHITE, rounded=True)
    add_rect(slide, Inches(0.55), Inches(1.4), Inches(6.0), Inches(0.62), NAVY)
    add_textbox(slide, Inches(0.8), Inches(1.52), Inches(5.5), Inches(0.4), "Kas / Bank", size=20, bold=True, color=WHITE)
    add_textbox(slide, Inches(0.8), Inches(2.25), Inches(5.5), Inches(2.3), "Di mana uang berada.\nContoh hari ini: DEMO-KAS dan DEMO-BCA.\nSaldo di form tidak bisa diketik. Angka datang dari buku besar.", size=16, color=SLATE)
    add_rect(slide, Inches(6.8), Inches(1.4), Inches(6.0), Inches(3.5), WHITE, rounded=True)
    add_rect(slide, Inches(6.8), Inches(1.4), Inches(6.0), Inches(0.62), TEAL)
    add_textbox(slide, Inches(7.05), Inches(1.52), Inches(5.5), Inches(0.4), "Dana Amanah", size=20, bold=True, color=WHITE)
    add_textbox(slide, Inches(7.05), Inches(2.25), Inches(5.5), Inches(2.3), "Untuk apa uang boleh dipakai.\nDEMO-ZKT dan DEMO-YATIM terikat.\nBiaya bank tidak boleh mengambil dari dana terikat.", size=16, color=SLATE)
    add_rect(slide, Inches(0.55), Inches(5.15), Inches(12.25), Inches(1.55), AMBER_BG, rounded=True)
    add_textbox(slide, Inches(0.85), Inches(5.35), Inches(11.7), Inches(1.15), "Total kas/bank harus sama dengan total Dana Amanah. Di dashboard, kondisi itu tampil sebagai Seimbang. Transaksi draft belum mengubah saldo. Saldo berubah saat transaksi disetujui.", size=16, color=SLATE)
    add_footer(slide, 9)


def stations_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Giliran tiap peran", "Kerjakan dari atas ke bawah")
    add_table(
        slide,
        ["Urutan", "Siapa", "Menu", "Selesai jika"],
        [
            ["1", "Semua pengurus", "Login", "Masuk dengan email sendiri"],
            ["2", "Erwin Ahadi", "Penerimaan", "Satu draft di-submit"],
            ["3", "Novianto", "Approval", "Penerimaan submitted disetujui"],
            ["4", "Erwin Ahadi", "Pengeluaran", "Satu draft di-submit"],
            ["5", "Syarif Ahmadi", "Approval", "Pengeluaran submitted terverifikasi"],
            ["6", "Johan Saputra", "Approval", "Pengeluaran verified disetujui"],
            ["7", "Novianto + Erwin", "Kas kecil", "Float draft dicairkan atau belanja dicatat"],
            ["8", "Firman + Johan", "Pengajuan Bantuan", "Satu kartu bergeser satu tahap"],
            ["9", "Laptop cadangan", "Portal Donatur", "donatur@sima.test hanya melihat donasinya"],
            ["10", "Muharrar", "Laporan", "Selisih rekonsiliasi global tetap 0"],
        ],
        Inches(0.4), Inches(1.28), Inches(12.5), Inches(5.55),
        font=12,
        widths=[Inches(1.15), Inches(2.7), Inches(2.7), Inches(5.95)],
    )
    add_footer(slide, 10)


def step_login(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Masuk dengan akun Anda", "Giliran 1  ·  semua orang")
    bullets(slide, [
        "Buka https://sima.ahpeduli.org/auth/v2/login",
        "Pakai email pengurus di slide akun. Password sementara: @Password123",
        "Setelah masuk, buka menu akun di kanan atas → Ubah password. Simpan password baru sendiri.",
        "Erwin menginput transaksi. Novianto menyetujui uang. Keduanya tidak memakai menu yang sama.",
        "Jika gagal, jangan membuat user baru. Minta fasilitator mengecek ejaan email dan peran.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 11)


def step_receipt(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Submit satu penerimaan", "Giliran 2  ·  Erwin Ahadi")
    bullets(slide, [
        "Menu Keuangan → Penerimaan.",
        "Data demo sudah punya draft, submitted, dan approved. Pilih yang masih draft.",
        "Pastikan total alokasi dana sama dengan nominal penerimaan.",
        "Dana yang dipakai: DEMO-ZKT, DEMO-YATIM, atau DEMO-UMUM. Jangan pilih dana sistem.",
        "Submit. Jangan approve dari akun ini. Serahkan ke Novianto.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 12)
    add_notes(slide, "Penerimaan approved menaikkan kas dan dana sekaligus. Sebelum approve, saldo belum berubah.")


def step_receipt_approve(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Setujui penerimaan", "Giliran 3  ·  Novianto")
    bullets(slide, [
        "Menu Approval. Buka penerimaan berstatus submitted.",
        "Cek donatur, rekening, dan alokasi. Total alokasi harus sama dengan nominal.",
        "Approve satu transaksi saja, supaya auditor masih bisa membedakan jejaknya.",
        "Kembali ke dashboard. Saldo rekening dan dana terkait naik.",
        "Jika ditolak, Erwin memperbaiki draft. Jangan buat penerimaan baru untuk saldo yang sudah ada.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 13)


def step_expense(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Submit satu pengeluaran", "Giliran 4  ·  Erwin Ahadi")
    bullets(slide, [
        "Menu Keuangan → Pengeluaran. Gunakan yang masih draft, atau buat satu nominal kecil.",
        "Pilih rekening DEMO-BCA atau DEMO-KAS.",
        "Isi sumber Dana Amanah. Jumlah sumber harus sama dengan nominal pengeluaran.",
        "Boleh lebih dari satu dana. Jangan mengambil dana restricted untuk biaya bank.",
        "Submit, lalu berhenti. Berikutnya giliran Syarif Ahmadi.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 14)


def step_verify(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Verifikasi pengeluaran", "Giliran 5  ·  Syarif Ahmadi")
    bullets(slide, [
        "Menu Approval. Antrian Anda adalah pengeluaran submitted.",
        "Cek penerima, nominal, lampiran, dan sumber dana.",
        "Verify jika lengkap. Status menjadi verified, saldo belum berubah.",
        "Tolak jika ada yang kurang, dan tulis alasannya.",
        "Anda tidak menyetujui dampak ke buku besar. Itu giliran Johan Saputra.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 15)


def step_expense_approve(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Setujui pengeluaran", "Giliran 6  ·  Johan Saputra")
    bullets(slide, [
        "Menu Approval. Ambil pengeluaran yang sudah verified.",
        "Approve hanya jika saldo rekening dan saldo dana mencukupi.",
        "Setelah approved, saldo turun. Transaksi tidak bisa diedit.",
        "Jika salah, jangan hapus. Pakai reversal dan tulis alasan.",
        "Coba sekali menolak approve saat sumber dana kurang, supaya tim melihat penjagaan sistem.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 16)


def step_petty(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Kas kecil", "Giliran 7  ·  Novianto mencairkan, Erwin mencatat")
    bullets(slide, [
        "Menu Kas Kecil. Uang muka hanya masuk lewat pengeluaran yang tertaut dan sudah approved.",
        "Novianto mencairkan float yang masih draft.",
        "Erwin mencatat belanja pada float yang sedang dilaporkan. Jika tombolnya tidak ada, float masih di akun contoh petugas.bantuan@sima.test.",
        "Baris belanja tidak membuat jurnal baru. Jurnal ada di pengeluaran pencairan.",
        "Jangan tutup float hari ini jika tim lain masih ingin melihat isinya.",
    ], top=1.45, size=17, gap=0.7)
    add_footer(slide, 17)
    add_notes(slide, "Kontrak lengkap ada di docs/KAS-KECIL.md. PJ adalah penugasan, bukan role terpisah. Di data demo, PJ-nya petugas bantuan.")


def step_grant(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Pengajuan bantuan", "Giliran 8  ·  Firman Ardiansyah dan Johan Saputra")
    bullets(slide, [
        "Menu Pengajuan Bantuan, papan /dashboard/bantuan.",
        "Firman membuat satu kartu perorangan, atau menggeser kartu yang sudah ada satu tahap.",
        "Johan menugaskan atau menyetujui. Jangan langsung menandai selesai.",
        "Selesai butuh pengeluaran tertaut yang sudah approved, plus bukti serah terima.",
        "Bantuan massal memakai satu kartu kolektif, bukan satu kartu per warga.",
    ], top=1.45, size=17, gap=0.7)
    add_footer(slide, 18)


def step_donor(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Portal donatur", "Giliran 9  ·  akun contoh, laptop terpisah")
    bullets(slide, [
        "Masuk dengan donatur@sima.test / password. Bukan email pengurus.",
        "Buka Portal Donatur.",
        "Yang tampil hanya penerimaan approved milik Budi Santoso.",
        "Penerimaan orang lain, draft, dan saldo lembaga tidak boleh terlihat.",
        "Jika donatur melihat menu keuangan, catat sebagai temuan.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 19)


def step_auditor(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Cek buku bersama", "Giliran 10  ·  Muharrar, dilihat semua")
    bullets(slide, [
        "Dashboard: status rekonsiliasi harus tetap Seimbang.",
        "Laporan → Saldo Dana Amanah dan Laporan → Ledger.",
        "Laporan → Saldo Awal: ada baris pembukaan untuk akun demo.",
        "Menu audit menampilkan siapa yang submit, verify, dan approve tadi.",
        "Jika selisih bukan nol, berhenti. Jangan tambah transaksi baru sampai penyebabnya ketemu.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 20)


def password_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Ganti password sementara", "Setiap pengurus, setelah berhasil masuk")
    add_rect(slide, Inches(0.55), Inches(1.4), Inches(12.2), Inches(1.7), AMBER_BG, rounded=True)
    add_textbox(
        slide, Inches(0.85), Inches(1.58), Inches(11.6), Inches(1.35),
        "Password awal @Password123 hanya untuk masuk pertama. Menu akun di kanan atas → Ubah password. Isi password saat ini, password baru, lalu ulangi password baru.",
        size=18, color=SLATE,
    )
    bullets(slide, [
        "Minimal 8 karakter, dan harus berbeda dari @Password123.",
        "Sesi di perangkat lain berakhir. Sesi yang dipakai untuk mengganti password tetap masuk.",
        "Jika lupa password baru, fasilitator mereset dari Pengaturan → detail pengguna.",
        "Akun contoh @sima.test tidak dipakai pengurus, dan tidak boleh dibawa ke produksi.",
    ], top=3.4, size=16, gap=0.58)
    add_footer(slide, 21)


def failures_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Yang memang harus gagal", "Coba satu, lalu catat pesan yang muncul")
    add_table(
        slide,
        ["Percobaan", "Hasil yang benar"],
        [
            ["Isi saldo langsung di form Kas/Bank", "Angka tidak bisa diketik"],
            ["Submit penerimaan tanpa alokasi yang cukup", "Ditolak"],
            ["Biaya bank dari DEMO-ZKT", "Ditolak, karena dana terikat"],
            ["Approve pengeluaran yang saldonya kurang", "Ditolak"],
            ["Hapus transaksi yang sudah approved", "Tidak ada jalan hapus; koreksi lewat reversal"],
            ["Donatur membuka daftar seluruh penerimaan", "Tidak punya menu itu"],
        ],
        Inches(0.55), Inches(1.4), Inches(12.2), Inches(5.2),
        font=14,
        widths=[Inches(6.3), Inches(5.9)],
    )
    add_footer(slide, 22)


def findings_slide(prs):
    slide = blank(prs)
    add_bg(slide)
    content_title(slide, "Catat temuan sebelum pulang", "Satu daftar bersama")
    bullets(slide, [
        "Siapa, peran, menu, dan apa yang diklik.",
        "Yang diharapkan terjadi, dan yang benar-benar terjadi.",
        "Cuplikan layar pesan error. Jangan hanya menulis “gagal”.",
        "Pisahkan: salah paham alur, bug, atau data demo yang habis terpakai.",
        "Admin tidak membuat user kantor di staging hari ini. Akun organisasi menyusul saat go-live.",
    ], top=1.45, size=18, gap=0.72)
    add_footer(slide, 23)


def close_slide(prs):
    slide = blank(prs)
    add_bg(slide, NAVY)
    add_rect(slide, 0, 0, Inches(0.22), H, GOLD)
    add_textbox(slide, Inches(0.85), Inches(1.5), Inches(11.5), Inches(0.35), "SELESAI JIKA", size=14, bold=True, color=GOLD)
    lines = [
        "Setiap peran sudah masuk dan menyelesaikan gilirannya.",
        "Ada satu penerimaan dan satu pengeluaran yang baru disetujui.",
        "Dashboard masih Seimbang.",
        "Temuan tercatat, termasuk yang sengaja gagal.",
        "Database staging ini tidak dipromosikan menjadi produksi.",
    ]
    for index, line in enumerate(lines):
        add_textbox(slide, Inches(0.85), Inches(2.15 + index * 0.7), Inches(11.2), Inches(0.55), f"{index + 1}    {line}", size=20, color=WHITE)
    add_textbox(slide, Inches(0.85), Inches(6.15), Inches(11.2), Inches(0.7), "Panduan lengkap: docs/PANDUAN-MULAI.md\nKontrak bantuan dan kas kecil: docs/BANTUAN.md  ·  docs/KAS-KECIL.md", size=14, color=RGBColor(0xCB, 0xD5, 0xE1))
    add_notes(slide, "Ingatkan lagi: akun demo dan password demo tinggal di staging.")


def build():
    presentation = Presentation()
    presentation.slide_width = W
    presentation.slide_height = H
    presentation.core_properties.title = "SIMA — Uji coba staging 23 Oktober 2026"
    presentation.core_properties.subject = "Urutan setup data demo dan uji coba tim"
    title_slide(presentation)
    rules_slide(presentation)
    url_slide(presentation)
    accounts_slide(presentation)
    agenda_slide(presentation)
    it_backup_slide(presentation)
    it_seed_slide(presentation)
    it_check_slide(presentation)
    concept_slide(presentation)
    stations_slide(presentation)
    step_login(presentation)
    step_receipt(presentation)
    step_receipt_approve(presentation)
    step_expense(presentation)
    step_verify(presentation)
    step_expense_approve(presentation)
    step_petty(presentation)
    step_grant(presentation)
    step_donor(presentation)
    step_auditor(presentation)
    password_slide(presentation)
    failures_slide(presentation)
    findings_slide(presentation)
    close_slide(presentation)
    if len(presentation.slides) != TOTAL:
        raise SystemExit(f"Jumlah slide {len(presentation.slides)} tidak sama dengan {TOTAL}")
    presentation.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
